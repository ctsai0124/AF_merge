import io
import os
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import paycheck
from app import app, build_cross_check_excel, build_roster_template_xlsx, _category_summary_text


def workbook_bytes(sheet_name, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, '') for header in headers])
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


class CrossCheckTemplateTests(unittest.TestCase):
    def test_output_uses_confirmed_official_template_structure(self):
        values = [
            1, '測試國小', '一般人員3人、職工1人、約聘雇人員2人',
            'A0001：3人', 100000, 'B10021：3人', 50000,
            'C1001：1人', 5000, '無', 0,
            100000, 50000, 5000, 0, '考績晉級差額，下月補發', 'N', '', '',
        ]
        workbook = openpyxl.load_workbook(build_cross_check_excel(values))
        sheet = workbook['互核結果表範本']

        self.assertEqual(sheet.max_row, 15)
        self.assertEqual(sheet.max_column, 19)
        self.assertEqual(sheet['A1'].value, '薪資互核結果表')
        self.assertEqual(sheet['B4'].value, '測試國小')
        self.assertEqual(sheet['P4'].value, '考績晉級差額，下月補發')
        self.assertEqual(sheet['Q4'].value, 'N')
        self.assertTrue(all(sheet.cell(5, col).value is None for col in range(1, 20)))
        self.assertEqual(
            {str(rng) for rng in sheet.merged_cells.ranges},
            {'A1:S1', 'A2:A3', 'B2:B3', 'C2:K2', 'L2:O2',
             'P2:P3', 'Q2:Q3', 'R2:R3', 'S2:S3', 'A6:S15'},
        )
        self.assertEqual(sheet.print_area, "'互核結果表範本'!$A$1:$S$15")
        self.assertEqual(sheet.page_setup.orientation, 'landscape')
        self.assertEqual(sheet.page_setup.paperSize, 9)
        self.assertEqual(sheet.page_setup.scale, 50)
        self.assertEqual(sheet.column_dimensions['P'].width, 62.125)
        self.assertEqual(sheet.row_dimensions[4].height, 178.65)
        self.assertIn('(一般人員、職工、約聘雇人員、政務人員)', sheet['A6'].value)
        self.assertNotIn('教育警察人員', sheet['A6'].value)

    def test_people_categories_are_only_broad_choices(self):
        self.assertEqual(
            paycheck.FINE_CATEGORY_OPTIONS,
            ['一般人員', '職工', '約聘雇人員', '政務人員'],
        )
        self.assertEqual(paycheck.category_from_salary_table('A00011'), '一般人員')
        self.assertEqual(paycheck.category_from_salary_table('A0004'), '約聘雇人員')

    def test_old_saved_fine_categories_remain_compatible(self):
        self.assertEqual(paycheck.rollup_category('教師'), '一般人員')
        self.assertEqual(paycheck.rollup_category('工友'), '職工')
        self.assertEqual(paycheck.rollup_category('約僱'), '約聘雇人員')
        self.assertEqual(paycheck.rollup_category('約聘僱人員'), '約聘雇人員')

    def test_summary_uses_broad_category_order(self):
        summary = {'職工': 1, '一般人員': 3, '約聘雇人員': 2}
        self.assertEqual(
            _category_summary_text(summary, paycheck.OFFICIAL_CATEGORIES),
            '一般人員3人、職工1人、約聘雇人員2人',
        )

    def test_fixed_roster_template_has_optional_category_dropdown(self):
        workbook = openpyxl.load_workbook(build_roster_template_xlsx())
        sheet = workbook['input']
        self.assertEqual(sheet['D1'].value, '人員種類')
        self.assertEqual(sheet['D2'].value, '一般人員')
        self.assertEqual(sheet['D3'].value, '職工')
        self.assertEqual(sheet['D4'].value, '約聘雇人員')
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        validation = sheet.data_validations.dataValidation[0]
        self.assertEqual(
            validation.formula1,
            '"一般人員,職工,約聘雇人員,政務人員"',
        )
        self.assertIn('D2:D500', str(validation.sqref))

    def test_difference_reason_is_written_to_formal_report(self):
        roster = workbook_bytes('input', [{
            '序號': 1, '姓名': '王小明', '身分證字號': 'A123456789',
            '人員種類': '一般人員',
        }])
        af = workbook_bytes('AF', [{
            '姓名': '王小明', '身分證字號': 'A123456789',
            '薪俸表別': 'A0001', '支領數額': 50000,
            '專業加給表別': 'B10021', '支領數額.1': 20000,
            '職務加給表別': '', '支領數額.2': 0,
            '地域加給表別': '', '支領數額.3': 0,
        }])
        with tempfile.TemporaryDirectory() as data_dir, patch.dict(
            os.environ, {'DATA_DIR': data_dir}
        ):
            response = app.test_client().post('/cross-check', data={
                'roster': (io.BytesIO(roster), '固定清冊.xlsx'),
                'af': (io.BytesIO(af), 'AF_000000000X_11509.xlsx'),
                'difference_reason': '考績晉級差額，下月補發',
            })

        self.assertEqual(response.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(response.data))['互核結果表範本']
        self.assertEqual(sheet['P4'].value, '考績晉級差額，下月補發')
        self.assertEqual(sheet['Q4'].value, 'N')


if __name__ == '__main__':
    unittest.main()
